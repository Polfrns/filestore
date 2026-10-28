"""Bounded attachment preparation and output validation; no user code is executed."""
from __future__ import annotations

import argparse
import gzip
import bz2
import lzma
import hashlib
import io
import json
import os
from pathlib import Path, PurePosixPath
import re
import shutil
import signal
import stat
import subprocess
import tarfile
import tempfile
import warnings
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import asdict, dataclass, field

MAX_FILE_BYTES = 15 * 1024 * 1024
MAX_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 500
MAX_TEXT_CHARS = 120_000
MAX_PREVIEW_PAGES = 4
IMAGE_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.gif', '.webp', '.bmp', '.tif', '.tiff', '.avif', '.heic', '.heif', '.svg'}
OFFICE_EXTENSIONS = {'.docx', '.xlsx', '.pptx', '.doc', '.xls', '.ppt', '.odt', '.ods', '.odp', '.rtf', '.docm', '.xlsm', '.pptm', '.xlsb'}
TEXT_EXTENSIONS = {'.txt', '.md', '.csv', '.tsv', '.json', '.jsonl', '.xml', '.yaml', '.yml', '.toml', '.html', '.htm', '.css', '.py', '.js', '.jsx', '.mjs', '.cjs', '.ts', '.tsx', '.c', '.cc', '.cpp', '.h', '.hpp', '.java', '.go', '.rs', '.sh', '.ps1', '.sql', '.log', '.ini', '.cfg', '.conf', '.tex', '.eml', '.srt', '.vtt', '.ipynb', '.rst', '.rb', '.php', '.kt', '.swift', '.r', '.bat', '.env', '.properties'}
MEDIA_EXTENSIONS = {'.wav', '.mp3', '.ogg', '.flac', '.m4a', '.aac', '.opus', '.mp4', '.mov', '.mkv', '.webm', '.avi'}


class FileToolError(Exception):
    def __init__(self, message: str, code: str = 'invalid_file') -> None:
        super().__init__(message)
        self.code = code


@dataclass
class PreparedFile:
    text: str = ''
    images: list[str] = field(default_factory=list)
    working_path: str | None = None
    archive_root: str | None = None
    warnings: list[str] = field(default_factory=list)


@dataclass
class ExpansionBudget:
    written: int = 0
    names: set[str] = field(default_factory=set)

    def member(self, name: str, size: int) -> PurePosixPath:
        parts = PurePosixPath(name.replace('\\', '/'))
        if (not name or name.startswith(('/', '\\')) or '..' in parts.parts or ':' in name
                or '\x00' in name or not parts.parts or len(parts.parts) > 20
                or any(p.rstrip(' .') != p or re.fullmatch(r'(?i)(con|prn|aux|nul|com[1-9]|lpt[1-9])(\..*)?', p) for p in parts.parts)):
            raise FileToolError('Unsafe archive path: ' + repr(name), 'unsafe_archive')
        canonical = str(parts).casefold()
        if canonical in self.names:
            raise FileToolError('Duplicate archive path: ' + repr(name), 'unsafe_archive')
        self.names.add(canonical)
        if len(self.names) > MAX_ARCHIVE_ENTRIES or size < 0 or size > MAX_FILE_BYTES:
            raise FileToolError('Archive exceeds file/count limits.', 'archive_limit')
        return parts

    def copy(self, source: io.BufferedIOBase, target: Path, expected: int | None = None) -> None:
        target.parent.mkdir(parents=True, exist_ok=True)
        count = 0
        with target.open('xb') as output:
            while chunk := source.read(65536):
                count += len(chunk)
                self.written += len(chunk)
                if count > MAX_FILE_BYTES or self.written > MAX_EXPANDED_BYTES:
                    raise FileToolError('Archive exceeds expanded-byte limits.', 'archive_limit')
                output.write(chunk)
        if expected is not None and count != expected:
            raise FileToolError('Archive member size mismatch.', 'invalid_archive')


def workspace_file(workspace: Path, name: str) -> Path:
    path = Path(name)
    candidate = path if path.is_absolute() else workspace / path
    if candidate.is_symlink():
        raise FileToolError('Symbolic links are not user files.', 'unsafe_path')
    resolved = candidate.resolve()
    if not resolved.is_relative_to(workspace) or resolved == workspace:
        raise FileToolError('Path is outside the request workspace.', 'unsafe_path')
    return resolved


def bounded_text(text: str) -> str:
    return text if len(text) <= MAX_TEXT_CHARS else text[:MAX_TEXT_CHARS] + '\n[Extracted preview truncated; read the original for remaining content.]'


def xml_root(data: bytes) -> ET.Element:
    if len(data) > MAX_FILE_BYTES or re.search(br'<!\s*(DOCTYPE|ENTITY)', data.replace(b'\x00', b''), re.I):
        raise FileToolError('Unsafe or oversized XML.', 'invalid_xml')
    return ET.fromstring(data)


def package_xml(path: Path, matcher: re.Pattern[str]) -> str:
    parts: list[str] = []
    total = 0
    with zipfile.ZipFile(path) as archive:
        if len(archive.infolist()) > 5000:
            raise FileToolError('Office package contains too many entries.')
        for item in archive.infolist():
            if not matcher.fullmatch(item.filename):
                continue
            total += item.file_size
            if item.file_size > MAX_FILE_BYTES or total > MAX_EXPANDED_BYTES:
                raise FileToolError('Office XML exceeds expanded-byte limits.')
            root = xml_root(archive.read(item))
            lines: list[str] = []
            for element in root.iter():
                local = element.tag.rsplit('}', 1)[-1]
                if local in {'t', 'p', 'h', 'span', 'description', 'title'} and element.text:
                    lines.append(element.text)
                elif local in {'tab', 'line-break', 'br'}:
                    lines.append('\t' if local == 'tab' else '\n')
            parts.append(' '.join(lines) if item.filename == 'word/document.xml' else item.filename + '\n' + ' '.join(lines))
            if sum(len(p) for p in parts) > MAX_TEXT_CHARS:
                break
    return bounded_text('\n\n'.join(parts))


def executable(name: str) -> str:
    configured = os.environ.get('BRIDGE_' + name.upper())
    found = configured or shutil.which(name)
    if name == 'soffice' and not found and os.name == 'nt':
        for base in (os.environ.get('PROGRAMFILES', ''), os.environ.get('PROGRAMFILES(X86)', '')):
            candidate = Path(base) / 'LibreOffice/program/soffice.exe'
            if candidate.is_file():
                found = str(candidate)
                break
    if not found:
        raise FileToolError(f'{name} is not installed. Install it on the bridge host or use the supplied Docker image.', 'missing_tool')
    return found


def run_tool(args: list[str], timeout: int = 40) -> str:
    with subprocess.Popen(args, stdout=subprocess.PIPE, stderr=subprocess.PIPE, start_new_session=os.name != 'nt') as process:
        try:
            stdout, stderr = process.communicate(timeout=timeout)
        finally:
            if os.name != 'nt':
                try:
                    os.killpg(process.pid, signal.SIGKILL)
                except ProcessLookupError:
                    pass
            elif process.poll() is None:
                subprocess.run(['taskkill', '/PID', str(process.pid), '/T', '/F'], capture_output=True, check=False)
            if process.poll() is None:
                process.kill()
                process.wait()
        if process.returncode:
            raise FileToolError('External file tool failed: ' + stderr.decode(errors='replace')[-1500:], 'tool_failed')
    if len(stdout) > 2 * 1024 * 1024:
        raise FileToolError('External tool output exceeds preview limit.')
    return stdout.decode(errors='replace')


def convert_office(source: Path, output: Path, workspace: Path, recalculate: bool = False) -> Path:
    if output.suffix.lower() not in OFFICE_EXTENSIONS | {'.pdf', '.txt', '.csv', '.html'}:
        raise FileToolError('Unsupported Office conversion target.', 'unsupported_conversion')
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='office-', dir=workspace) as folder:
        temp = Path(folder)
        profile = temp / 'profile' / 'user'
        profile.mkdir(parents=True)
        # Fresh profiles have no trusted sources; disable automatic document macros.
        # Calc normally keeps cached results on load; mode 0 recalculates every formula.
        recalc = ''.join(f'<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="{name}" oor:op="fuse"><value>0</value></prop></item>'
                         for name in ('OOXMLRecalcMode', 'ODFRecalcMode')) if recalculate else ''
        (profile / 'registrymodifications.xcu').write_text('<oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>' + recalc + '</oor:items>', encoding='utf-8')
        run_tool([executable('soffice'), '-env:UserInstallation=' + (temp / 'profile').as_uri(), '--headless', '--convert-to', output.suffix[1:], '--outdir', str(temp), str(source)])
        converted = temp / (source.stem + output.suffix)
        if not converted.is_file() or converted.stat().st_size == 0:
            raise FileToolError('LibreOffice did not produce the requested conversion.', 'unsupported_conversion')
        shutil.copyfile(converted, output)
    return output


def image_preview(source: Path, destination: Path) -> str:
    from PIL import Image, ImageOps
    Image.MAX_IMAGE_PIXELS = 40_000_000
    try:
        from pillow_heif import register_heif_opener
        register_heif_opener()
    except ImportError:
        if source.suffix.lower() in {'.heic', '.heif'}:
            raise FileToolError('HEIC/HEIF support requires pillow-heif; rerun setup:files.', 'missing_tool')
    destination.parent.mkdir(parents=True, exist_ok=True)
    with warnings.catch_warnings():
        warnings.simplefilter('error', Image.DecompressionBombWarning)
        if source.suffix.lower() == '.svg':
            root = xml_root(source.read_bytes())
            for element in root.iter():
                if any(key.rsplit('}', 1)[-1] == 'href' and not value.startswith(('#', 'data:')) for key, value in element.attrib.items()):
                    raise FileToolError('SVG external resources are not allowed.', 'unsafe_image')
            import pymupdf
            with pymupdf.open(source) as doc:
                page = doc[0]
                scale = min(2, 1600 / max(page.rect.width, page.rect.height))
                page.get_pixmap(matrix=pymupdf.Matrix(scale, scale)).save(destination)
        else:
            with Image.open(source) as opened:
                image = ImageOps.exif_transpose(opened).convert('RGBA')
                image.thumbnail((1600, 1600))
                image.save(destination, 'PNG')
    return str(destination)


def extract_archive(source: Path, destination: Path) -> list[Path]:
    destination.mkdir(parents=True, exist_ok=False)
    budget = ExpansionBudget()
    extracted: list[Path] = []
    if source.suffix.lower() == '.zip':
        with zipfile.ZipFile(source) as archive:
            for item in archive.infolist():
                member = budget.member(item.filename, item.file_size)
                if item.flag_bits & 1 or stat.S_ISLNK(item.external_attr >> 16):
                    raise FileToolError('Encrypted files and links are not extracted.', 'unsafe_archive')
                target = destination.joinpath(*member.parts)
                if item.is_dir():
                    target.mkdir(parents=True, exist_ok=True)
                else:
                    with archive.open(item) as stream:
                        budget.copy(stream, target, item.file_size)
                    extracted.append(target)
    elif source.suffix.lower() == '.7z':
        import py7zr
        class LimitedIO(py7zr.Py7zIO):
            def __init__(self) -> None:
                self.buffer = io.BytesIO()
            def write(self, data: bytes) -> int:
                budget.written += len(data)
                if self.buffer.tell() + len(data) > MAX_FILE_BYTES or budget.written > MAX_EXPANDED_BYTES:
                    raise FileToolError('7z exceeds expanded-byte limits.', 'archive_limit')
                return self.buffer.write(data)
            def read(self, size: int | None = None) -> bytes:
                return self.buffer.read(-1 if size is None else size)
            def seek(self, offset: int, whence: int = 0) -> int:
                return self.buffer.seek(offset, whence)
            def flush(self) -> None:
                self.buffer.flush()
            def size(self) -> int:
                return len(self.buffer.getbuffer())
        class LimitedFactory(py7zr.WriterFactory):
            def __init__(self) -> None:
                self.files: dict[str, LimitedIO] = {}
            def create(self, filename: str) -> LimitedIO:
                local = workspace_file(destination.resolve(), filename).relative_to(destination.resolve()).as_posix()
                if local.casefold() not in budget.names:
                    raise FileToolError('Unexpected 7z extraction path.', 'unsafe_archive')
                self.files[local] = LimitedIO()
                return self.files[local]
        factory = LimitedFactory()
        with py7zr.SevenZipFile(source) as archive:
            for item in archive.list():
                budget.member(item.filename, item.uncompressed or 0)
            if archive.needs_password():
                raise FileToolError('Encrypted archives require a separate decrypted upload.', 'unsafe_archive')
            for entry in archive.files:
                if entry.is_symlink or entry.is_junction:
                    raise FileToolError('7z links are not extracted.', 'unsafe_archive')
            archive.extractall(path=destination, factory=factory)
        for name, stream in factory.files.items():
            target = destination.joinpath(*PurePosixPath(name.replace('\\', '/')).parts)
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(stream.buffer.getvalue())
            extracted.append(target)
    elif tarfile.is_tarfile(source):
        with tarfile.open(source, 'r:*') as archive:
            for item in archive:
                member = budget.member(item.name, item.size)
                target = destination.joinpath(*member.parts)
                if item.isdir():
                    target.mkdir(parents=True, exist_ok=True)
                elif item.isfile():
                    with archive.extractfile(item) as stream:
                        budget.copy(stream, target, item.size)
                    extracted.append(target)
                else:
                    raise FileToolError('Archive links/devices are not extracted.', 'unsafe_archive')
    elif source.suffix.lower() in {'.gz', '.bz2', '.xz'}:
        member = budget.member(source.stem, 0)
        target = destination.joinpath(*member.parts)
        opener = {'.gz': gzip.open, '.bz2': bz2.open, '.xz': lzma.open}[source.suffix.lower()]
        with opener(source, 'rb') as stream:
            budget.copy(stream, target)
        extracted.append(target)
    else:
        raise FileToolError('Not a supported ZIP, TAR, gzip or 7z archive.', 'invalid_archive')
    return extracted


def check_office_package(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if (len(entries) > 5000 or sum(i.file_size for i in entries) > MAX_EXPANDED_BYTES
                or any(i.file_size > MAX_FILE_BYTES or i.flag_bits & 1 for i in entries)):
            raise FileToolError('Office package exceeds limits or is encrypted.')


def prepare_file(source: Path, workspace: Path) -> PreparedFile:
    suffix = source.suffix.lower()
    identity = hashlib.sha256(str(source.relative_to(workspace)).encode()).hexdigest()[:10]
    destination = workspace / '.bridge-prepared' / identity
    destination.mkdir(parents=True, exist_ok=True)
    result = PreparedFile(working_path=str(source))
    if suffix in {'.docx', '.docm', '.xlsx', '.xlsm', '.pptx', '.pptm', '.odt', '.ods', '.odp'}:
        check_office_package(source)
    if suffix in {'.zip', '.7z', '.tar', '.tgz', '.gz', '.bz2', '.xz'}:
        root = destination / 'extracted'
        files = extract_archive(source, root)
        result.archive_root = str(root)
        parts = ['Extracted archive; nested archives are retained without automatic recursion:']
        for path in files:
            parts.append(str(path.relative_to(root)))
            if path.suffix.lower() in TEXT_EXTENSIONS and sum(len(p) for p in parts) < MAX_TEXT_CHARS:
                parts.append(path.read_bytes()[:32000].decode('utf-8-sig', errors='replace'))
        for path in [p for p in files if p.suffix.lower() in IMAGE_EXTENSIONS | OFFICE_EXTENSIONS | {'.pdf'}][:8]:
            try:
                child = prepare_file(path, workspace)
                parts.append(str(path.relative_to(root)) + '\n' + child.text)
                result.images.extend(child.images[:max(0, 12 - len(result.images))])
                result.warnings.extend(child.warnings)
            except Exception as error:
                result.warnings.append(f'{path.relative_to(root)}: {error}')
        result.text = bounded_text('\n'.join(parts))
    elif suffix == '.pdf':
        import pymupdf
        parts = []
        with pymupdf.open(source) as doc:
            if doc.needs_pass:
                raise FileToolError('PDF is password protected.')
            for index in range(min(len(doc), 30)):
                page = doc[index]
                text = page.get_text()
                parts.append(f'Page {index + 1}:\n{text}')
                if index < MAX_PREVIEW_PAGES:
                    preview = destination / f'page-{index + 1}.png'
                    scale = min(2, 1600 / max(page.rect.width, page.rect.height))
                    page.get_pixmap(matrix=pymupdf.Matrix(scale, scale)).save(preview)
                    result.images.append(str(preview))
                    if not text.strip():
                        try:
                            parts.append(run_tool([executable('tesseract'), str(preview), 'stdout', '-l', os.environ.get('BRIDGE_OCR_LANG', 'eng+spa')]))
                        except FileToolError as error:
                            result.warnings.append(str(error))
            if len(doc) > 30:
                result.warnings.append('Only the first 30 pages are included in the text preview; read the original for the rest.')
        result.text = bounded_text('\n'.join(parts))
    elif suffix in IMAGE_EXTENSIONS:
        preview = destination / 'preview.png'
        result.images.append(image_preview(source, preview))
        result.text = 'Image preview uses the first frame/page; original resolution, animation and metadata remain in the original file.'
    elif suffix in {'.docx', '.docm'}:
        result.text = package_xml(source, re.compile(r'word/(document|header\d+|footer\d+|comments|footnotes|endnotes)\.xml'))
    elif suffix in {'.pptx', '.pptm'}:
        result.text = package_xml(source, re.compile(r'ppt/(slides/slide\d+|notesSlides/notesSlide\d+)\.xml'))
    elif suffix in {'.odt', '.ods', '.odp'}:
        result.text = package_xml(source, re.compile(r'(content|meta)\.xml'))
    elif suffix in {'.xlsx', '.xlsm'}:
        import openpyxl
        rows: list[str] = []
        book = openpyxl.load_workbook(source, read_only=True, data_only=False)
        try:
            chars = 0
            for sheet in list(book)[:20]:
                rows.append('Sheet: ' + sheet.title)
                for row in sheet.iter_rows(max_row=min(sheet.max_row or 1, 5000), max_col=min(sheet.max_column or 1, 100)):
                    value = '\t'.join('' if c.value is None else str(c.value) for c in row)
                    rows.append(value)
                    chars += len(value) + 1
                    if chars >= MAX_TEXT_CHARS:
                        break
                if chars >= MAX_TEXT_CHARS:
                    break
        finally:
            book.close()
        result.text = bounded_text('\n'.join(rows))
        result.warnings.append('Sheet preview is bounded to 20 sheets, 5000 rows and 100 columns per sheet; formulas are shown, not recalculated.')
    elif suffix in {'.doc', '.xls', '.ppt', '.rtf', '.xlsb'}:
        extension = {'.doc': '.docx', '.rtf': '.docx', '.xls': '.xlsx', '.xlsb': '.xlsx', '.ppt': '.pptx'}[suffix]
        converted = convert_office(source, destination / (source.stem + extension), workspace)
        converted_result = prepare_file(converted, workspace)
        result.text = converted_result.text
        result.working_path = str(converted)
        result.warnings.append('Converted working copy. Preserve the original extension by converting the edited file back with file-tools.py convert. Conversion may change layout or unsupported features.')
    elif suffix in MEDIA_EXTENSIONS:
        metadata = run_tool([executable('ffprobe'), '-v', 'error', '-protocol_whitelist', 'file,pipe', '-show_format', '-show_streams', '-of', 'json', str(source)])
        result.text = bounded_text(metadata)
        info = json.loads(metadata)
        duration = min(float(info.get('format', {}).get('duration', 0)), 3600)
        if any(stream.get('codec_type') == 'video' for stream in info.get('streams', [])):
            result.warnings.append('Only three sampled video frames are supplied; they do not describe the entire video.')
            for index in range(3):
                frame = destination / f'frame-{index + 1}.png'
                run_tool([executable('ffmpeg'), '-v', 'error', '-protocol_whitelist', 'file,pipe', '-ss', str(duration * index / 3), '-i', str(source), '-frames:v', '1', '-vf', "scale='min(1600,iw)':-2", '-y', str(frame)])
                if frame.is_file():
                    result.images.append(str(frame))
        if os.environ.get('BRIDGE_WHISPER_MODEL'):
            result.text += '\nTranscript:\n' + transcribe(source)
        else:
            result.warnings.append('Audio is not transcribed. Configure a provisioned local BRIDGE_WHISPER_MODEL and install file-requirements-media.txt for speech understanding. FFmpeg editing/conversion is available.')
    elif suffix == '.eml':
        from email import policy
        from email.parser import BytesParser
        message = BytesParser(policy=policy.default).parsebytes(source.read_bytes())
        body = message.get_body(preferencelist=('plain', 'html'))
        result.text = bounded_text(f'Subject: {message["Subject"]}\nFrom: {message["From"]}\nTo: {message["To"]}\n' + (body.get_content() if body else 'No readable message body.'))
    elif suffix in TEXT_EXTENSIONS or re.fullmatch(r'(?:\d+-)?(dockerfile|makefile)', source.name, re.I):
        data = source.read_bytes()
        encoding = 'utf-16' if data[:2] in {b'\xff\xfe', b'\xfe\xff'} else 'utf-8-sig'
        result.text = bounded_text(data.decode(encoding, errors='replace'))
    else:
        result.warnings.append('No format reader registered. The original binary is available; use a compatible installed tool or declare unsupported.')
    if suffix in {'.docx', '.docm', '.pptx', '.pptm', '.odt', '.odp'}:
        try:
            layout = convert_office(source, destination / 'layout.pdf', workspace)
            rendered = prepare_file(layout, workspace)
            result.images.extend(rendered.images)
            result.warnings.extend(rendered.warnings)
            result.warnings.append('Office layout preview is rendered by LibreOffice; it may differ from Microsoft Office. Inspect the original and preserve its styles.')
        except Exception as error:
            result.warnings.append('Office layout preview unavailable: ' + str(error))
    return result


def transcribe(source: Path) -> str:
    from faster_whisper import WhisperModel
    model_path = Path(os.environ.get('BRIDGE_WHISPER_MODEL', ''))
    if not model_path.is_dir():
        raise FileToolError('BRIDGE_WHISPER_MODEL must point to an already downloaded local model.', 'missing_tool')
    model = WhisperModel(str(model_path), device='cpu', compute_type='int8', local_files_only=True)
    segments, _ = model.transcribe(str(source), beam_size=1)
    text = ''
    for segment in segments:
        text += f'[{segment.start:.1f}-{segment.end:.1f}] {segment.text}\n'
        if len(text) > MAX_TEXT_CHARS:
            break
    return bounded_text(text)


FORMULA_ERRORS = {'#REF!', '#DIV/0!', '#NAME?', '#VALUE!', '#N/A', '#NUM!', '#NULL!', '#SPILL!', '#CALC!'}
MAX_FORMULA_CELLS = 2_000_000


def recalculate_workbook(source: Path, output: Path, workspace: Path) -> dict:
    """LibreOffice recalculates a copy; the result lists cells whose fresh value is an error."""
    if source.suffix.lower() not in {'.xlsx', '.xlsm', '.xls', '.xlsb', '.ods', '.csv'} or output.suffix.lower() != '.xlsx':
        raise FileToolError('Recalculation reads a spreadsheet and writes an .xlsx copy.', 'unsupported_conversion')
    convert_office(source, output, workspace, recalculate=True)
    import openpyxl
    errors: list[dict] = []
    count = cells = 0
    book = openpyxl.load_workbook(output, read_only=True, data_only=True)
    try:
        for sheet in book:
            for row in sheet.iter_rows():
                for cell in row:
                    cells += 1
                    if cells > MAX_FORMULA_CELLS:
                        raise FileToolError('Workbook exceeds the recalculation cell limit.', 'archive_limit')
                    if isinstance(cell.value, str) and cell.value in FORMULA_ERRORS:
                        count += 1
                        if len(errors) < 50:
                            errors.append({'sheet': sheet.title, 'cell': cell.coordinate, 'value': cell.value})
    finally:
        book.close()
    return {'path': str(output), 'error_count': count, 'errors': errors}


STRUCTURED_TEXT = {'.json', '.ipynb', '.jsonl', '.csv', '.tsv', '.py', '.toml'}


def read_text(path: Path, legacy: bool = False) -> str:
    data = path.read_bytes()
    encoding = 'utf-16' if data[:2] in {b'\xff\xfe', b'\xfe\xff'} else 'utf-8-sig'
    try:
        return data.decode(encoding)
    except UnicodeDecodeError as error:
        # Spreadsheet exports are often Windows-1252; preserving that encoding is valid.
        if legacy:
            return data.decode('latin-1')
        raise FileToolError(f'Text deliverable is not valid {encoding.upper()}: {error}.') from None


def validate_text(path: Path, suffix: str) -> None:
    text = read_text(path, legacy=suffix in {'.csv', '.tsv'})
    if suffix in {'.json', '.ipynb'}:
        json.loads(text)
    elif suffix == '.jsonl':
        for number, line in enumerate(text.splitlines(), 1):
            if line.strip():
                try:
                    json.loads(line)
                except json.JSONDecodeError as error:
                    raise FileToolError(f'Invalid JSON on line {number}: {error.msg}.') from None
    elif suffix in {'.csv', '.tsv'}:
        import csv
        try:
            for _ in csv.reader(io.StringIO(text, newline=''), delimiter='\t' if suffix == '.tsv' else ',', strict=True):
                pass
        except csv.Error as error:
            raise FileToolError(f'Malformed {suffix[1:].upper()}: {error}.') from None
    elif suffix == '.py':
        import ast
        # Parsing only: the deliverable is never compiled to bytecode or executed.
        try:
            ast.parse(text, filename=path.name)
        except SyntaxError as error:
            raise FileToolError(f'Python syntax error on line {error.lineno}: {error.msg}.') from None
    elif suffix == '.toml':
        try:
            import tomllib
        except ImportError:  # Python 3.10 has no TOML reader.
            return
        try:
            tomllib.loads(text)
        except tomllib.TOMLDecodeError as error:
            raise FileToolError(f'Invalid TOML: {error}.') from None


def validate_file(path: Path, workspace: Path) -> None:
    if not path.is_file() or not 0 < path.stat().st_size <= MAX_FILE_BYTES:
        raise FileToolError('Missing, empty or oversized deliverable.')
    suffix = path.suffix.lower()
    if suffix == '.pdf':
        import pymupdf
        with pymupdf.open(path) as doc:
            if doc.needs_pass or len(doc) == 0:
                raise FileToolError('Unreadable PDF.')
            doc[0].get_text()
    elif suffix in IMAGE_EXTENSIONS:
        with tempfile.TemporaryDirectory(dir=workspace) as folder:
            image_preview(path, Path(folder) / 'checked.png')
    elif suffix in {'.docx', '.docm', '.xlsx', '.xlsm', '.pptx', '.pptm', '.odt', '.ods', '.odp'}:
        required = 'content.xml' if suffix in {'.odt', '.ods', '.odp'} else {'.docx': 'word/document.xml', '.docm': 'word/document.xml', '.xlsx': 'xl/workbook.xml', '.xlsm': 'xl/workbook.xml', '.pptx': 'ppt/presentation.xml', '.pptm': 'ppt/presentation.xml'}[suffix]
        check_office_package(path)
        with zipfile.ZipFile(path) as archive:
            if archive.testzip() is not None:
                raise FileToolError('Corrupt Office package.')
            xml_root(archive.read(required))
        if suffix in {'.xlsx', '.xlsm'}:
            import openpyxl
            book = openpyxl.load_workbook(path, read_only=True)
            book.close()
    elif suffix in {'.zip', '.7z', '.tar', '.tgz', '.gz', '.bz2', '.xz'}:
        with tempfile.TemporaryDirectory(dir=workspace) as folder:
            extract_archive(path, Path(folder) / 'checked')
    elif suffix in {'.doc', '.xls', '.ppt', '.rtf', '.xlsb'}:
        with tempfile.TemporaryDirectory(dir=workspace) as folder:
            convert_office(path, Path(folder) / 'checked.pdf', workspace)
    elif suffix in MEDIA_EXTENSIONS:
        info = json.loads(run_tool([executable('ffprobe'), '-v', 'error', '-protocol_whitelist', 'file,pipe', '-show_streams', '-of', 'json', str(path)]))
        if not info.get('streams'):
            raise FileToolError('Media output has no readable streams.')
    elif suffix in {'.xml', '.svg'}:
        xml_root(path.read_bytes())
    elif suffix in STRUCTURED_TEXT:
        validate_text(path, suffix)


def main() -> None:
    def cancel_reader(_signum: int, _frame: object) -> None:
        raise SystemExit(1)
    signal.signal(signal.SIGTERM, cancel_reader)
    parser = argparse.ArgumentParser()
    parser.add_argument('operation', choices=['prepare', 'validate', 'convert', 'recalc', 'transcribe'])
    parser.add_argument('--workspace', required=True)
    parser.add_argument('--path', required=True)
    parser.add_argument('--output')
    args = parser.parse_args()
    workspace = Path(args.workspace).resolve(strict=True)
    try:
        source = workspace_file(workspace, args.path)
        if args.operation == 'prepare':
            result = asdict(prepare_file(source, workspace))
        elif args.operation == 'validate':
            validate_file(source, workspace)
            result = {'valid': True}
        elif args.operation == 'convert':
            if not args.output:
                raise FileToolError('Conversion requires --output.')
            target = workspace_file(workspace, args.output)
            convert_office(source, target, workspace)
            validate_file(target, workspace)
            result = {'path': str(target)}
        elif args.operation == 'recalc':
            if not args.output:
                raise FileToolError('Recalculation requires --output.')
            result = recalculate_workbook(source, workspace_file(workspace, args.output), workspace)
        else:
            result = {'text': transcribe(source)}
        print(json.dumps(result, ensure_ascii=False))
    except Exception as error:
        code = 'missing_tool' if isinstance(error, ImportError) else getattr(error, 'code', 'invalid_file')
        print(json.dumps({'error': str(error), 'code': code}))
        raise SystemExit(1)


if __name__ == '__main__':
    main()
